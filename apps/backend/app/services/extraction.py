"""Deterministic local text extraction for the first usable document library."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document as WordDocument
from openpyxl import load_workbook
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = (".csv", ".docx", ".md", ".pdf", ".txt", ".xlsm", ".xlsx")
MAX_EXTRACTED_CHARACTERS = 6_000_000
MAX_SPREADSHEET_CELLS = 500_000

MIME_TYPES = {
    ".csv": "text/csv",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".md": "text/markdown",
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}


@dataclass(frozen=True)
class ExtractionResult:
    text: str
    page_count: int | None
    word_count: int
    mime_type: str


def _bounded(text: str) -> str:
    normalized = text.replace("\x00", "").strip()
    if len(normalized) > MAX_EXTRACTED_CHARACTERS:
        return normalized[:MAX_EXTRACTED_CHARACTERS]
    return normalized


def _count_words(text: str) -> int:
    return len(re.findall(r"\w+", text, flags=re.UNICODE))


def _extract_pdf(path: Path) -> tuple[str, int]:
    reader = PdfReader(path, strict=False)
    pages: list[str] = []
    for page_number, page in enumerate(reader.pages, start=1):
        content = page.extract_text() or ""
        if content.strip():
            pages.append(f"[Page {page_number}]\n{content}")
    return "\n\n".join(pages), len(reader.pages)


def _extract_docx(path: Path) -> str:
    document = WordDocument(str(path))
    blocks = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            rows.append("\t".join(cell.text.strip() for cell in row.cells))
        if rows:
            blocks.append(f"[Table {table_index}]\n" + "\n".join(rows))
    return "\n\n".join(blocks)


def _extract_workbook(path: Path) -> str:
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    blocks: list[str] = []
    processed_cells = 0
    try:
        for worksheet in workbook.worksheets:
            rows: list[str] = []
            for row in worksheet.iter_rows(values_only=True):
                processed_cells += len(row)
                if processed_cells > MAX_SPREADSHEET_CELLS:
                    rows.append("[Spreadsheet extraction limit reached]")
                    break
                values = ["" if value is None else str(value) for value in row]
                if any(value for value in values):
                    rows.append("\t".join(values))
            blocks.append(f"[Sheet: {worksheet.title}]\n" + "\n".join(rows))
            if processed_cells > MAX_SPREADSHEET_CELLS:
                break
    finally:
        workbook.close()
    return "\n\n".join(blocks)


def _extract_text_file(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeError:
            continue
    return path.read_text(encoding="utf-8", errors="replace")


def _extract_csv(path: Path) -> str:
    raw = _extract_text_file(path)
    try:
        rows = csv.reader(raw.splitlines())
        return "\n".join("\t".join(row) for row in rows)
    except csv.Error:
        return raw


def extract_document(path: Path) -> ExtractionResult:
    extension = path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported document type: {extension or 'no extension'}")

    page_count: int | None = None
    if extension == ".pdf":
        text, page_count = _extract_pdf(path)
    elif extension == ".docx":
        text = _extract_docx(path)
    elif extension in {".xlsx", ".xlsm"}:
        text = _extract_workbook(path)
    elif extension == ".csv":
        text = _extract_csv(path)
    else:
        text = _extract_text_file(path)

    text = _bounded(text)
    return ExtractionResult(
        text=text,
        page_count=page_count,
        word_count=_count_words(text),
        mime_type=MIME_TYPES[extension],
    )
