from __future__ import annotations

from pathlib import Path

from app.services.extraction import extract_document
from docx import Document as WordDocument
from openpyxl import Workbook
from pypdf import PdfWriter


def test_extracts_word_tables_and_spreadsheet_cells(tmp_path: Path) -> None:
    word_path = tmp_path / "submittal.docx"
    word = WordDocument()
    word.add_paragraph("Concrete durability requirements")
    table = word.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Exposure class"
    table.cell(0, 1).text = "XS3"
    word.save(word_path)

    workbook_path = tmp_path / "schedule.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Materials"
    sheet.append(["Item", "Status"])
    sheet.append(["Waterproofing", "Approved with comments"])
    workbook.save(workbook_path)
    workbook.close()

    word_result = extract_document(word_path)
    workbook_result = extract_document(workbook_path)

    assert "Concrete durability requirements" in word_result.text
    assert "Exposure class\tXS3" in word_result.text
    assert "[Sheet: Materials]" in workbook_result.text
    assert "Waterproofing\tApproved with comments" in workbook_result.text


def test_blank_pdf_is_detected_as_having_no_extractable_text(tmp_path: Path) -> None:
    pdf_path = tmp_path / "scan.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    with pdf_path.open("wb") as stream:
        writer.write(stream)

    result = extract_document(pdf_path)

    assert result.page_count == 1
    assert result.text == ""
    assert result.word_count == 0
